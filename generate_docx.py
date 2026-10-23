import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set internal padding for a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def create_callout_box(doc, text_paragraphs, border_color="1E3A8A", bg_color="F0F4F8"):
    """Creates a visually appealing callout box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (thick bar)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    # Fill text
    p = cell.paragraphs[0]
    p.text = text_paragraphs[0][0]
    p.runs[0].font.bold = text_paragraphs[0][1].get("bold", False)
    p.runs[0].font.italic = text_paragraphs[0][1].get("italic", False)
    p.runs[0].font.size = Pt(text_paragraphs[0][1].get("size", 10.5))
    p.runs[0].font.color.rgb = text_paragraphs[0][1].get("color", RGBColor(30, 41, 59))
    p.paragraph_format.space_after = Pt(4)
    
    for text, style in text_paragraphs[1:]:
        p2 = cell.add_paragraph()
        run = p2.add_run(text)
        run.font.bold = style.get("bold", False)
        run.font.italic = style.get("italic", False)
        run.font.size = Pt(style.get("size", 10))
        run.font.color.rgb = style.get("color", RGBColor(51, 65, 85))
        p2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph() # Spacing after table

def style_table_header(row, col_widths, titles, bg_color="1E3A8A"):
    for i, title in enumerate(titles):
        cell = row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.text = title
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

def format_table_row(row, col_widths, values, is_even=False):
    bg_color = "F8FAFC" if is_even else "FFFFFF"
    for i, val in enumerate(values):
        cell = row.cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.text = str(val)
        run = p.runs[0]
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(30, 41, 59)

def generate_report():
    doc = Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    navy = RGBColor(30, 58, 138)       # Primary #1E3A8A
    slate_dark = RGBColor(30, 41, 59)  # Text #1E293B
    accent_blue = RGBColor(37, 99, 235)# Accent #2563EB

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("FlowTrack: Enterprise Visual Workflow & Project Management Platform")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = navy

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Comprehensive Technical Dossier, Architectural Analysis & Strategic Agenda")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = accent_blue

    # Meta Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.columns[0].width = Inches(3.25)
    meta_table.columns[1].width = Inches(3.25)
    
    meta_data = [
        [("Project Name", "FlowTrack Monorepo (@flowtrack)"), ("Version", "v0.3.0 (Enterprise Canvas Phase)")],
        [("Repository Root", "Project-managment-app"), ("Target Audience", "Project Managers, Tech Leads, Devs, & Executives")]
    ]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx in range(2):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            lbl, val = meta_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{lbl}: ")
            r1.font.bold = True
            r1.font.size = Pt(9)
            r1.font.color.rgb = navy
            r2 = p.add_run(val)
            r2.font.size = Pt(9)
            r2.font.color.rgb = slate_dark
    
    doc.add_paragraph()

    # SECTION 1: EXECUTIVE SUMMARY & MAIN AGENDA
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Executive Summary & Main Agenda")
    h1_run.font.color.rgb = navy
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    create_callout_box(doc, [
        ("THE MAIN AGENDA OF FLOWTRACK", {"bold": True, "size": 11, "color": navy}),
        ("FlowTrack is designed to solve a fundamental crisis in modern team operations: traditional project tools (Jira, Trello, Asana) force teams into rigid, linear tabular columns that hide dependencies, while diagramming software (Miro, Lucidchart) offers visual freedom but lacks execution rigor, data persistence, and audited accountability.", {"bold": False, "size": 10}),
        ("FlowTrack bridges this gap by providing an interactive, infinite-canvas workflow orchestrator (inspired by n8n, Automa, and React Flow) paired with enterprise-grade project governance: weighted mathematical progress rollups, strict multi-tier role-based permissions (RBAC), multi-level sub-canvases, audited deadline extensions, voice/photo remarks, document locking, and executive dossier reporting.", {"bold": False, "size": 10})
    ], border_color="1E3A8A", bg_color="EFF6FF")

    doc.add_paragraph(
        "Modern enterprise projects—whether in software engineering, construction, manufacturing, or cross-functional business initiatives—are rarely simple linear checklists. They involve complex prerequisites, conditional approval gates, interconnected dependencies, and multi-tier hierarchies. FlowTrack was conceived and engineered to serve five strategic pillars:"
    )

    pillars = [
        ("1. Visual Workflow Orchestration", "Treats tasks not as isolated tickets, but as nodes in a dynamic, directed graph. Teams visualize prerequisites, parallel execution streams, and critical paths at a single glance with 4-directional connection handles and custom sequence edges."),
        ("2. Real-World Accountability & Auditability", "Solves deadline slippage by forbidding silent date modifications. Whenever a milestone is exceeded or adjusted, the system triggers a formal Deadline Extension Protocol requiring an audited justification and supporting documentation."),
        ("3. Hierarchical Clarity Through Sub-Canvases", "Prevents canvas clutter by enabling infinite hierarchical drill-downs. High-level project stages encapsulate entire sub-canvases with nested nodes and breadcrumb navigation."),
        ("4. Deterministic Mathematical Rollups", "Eliminates subjective progress estimation. Group and project progress are mathematically computed using a weighted rollup engine that automatically factors task weights, blocked statuses, and overdue timestamps."),
        ("5. Enterprise Governance & Security", "Implements strict role separation across Super Admin, Manager, Team Lead, and Member roles, complete with confidential task locking, approval gates, and multi-language support (English, Hindi, Kannada).")
    ]
    for title, desc in pillars:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r_bold = p.add_run(f"{title}: ")
        r_bold.font.bold = True
        r_bold.font.color.rgb = navy
        r_desc = p.add_run(desc)
        r_desc.font.color.rgb = slate_dark

    doc.add_paragraph()

    # SECTION 2: SYSTEM ARCHITECTURE & MONOREPO STRUCTURE
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. System Architecture & Tech Stack")
    h2_run.font.color.rgb = navy
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "FlowTrack is engineered as a production-grade TypeScript monorepo using pnpm workspaces. The codebase cleanly separates presentation, business logic, persistence, and domain calculation engines into dedicated packages:"
    )

    arch_table = doc.add_table(rows=5, cols=3)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [1.5, 2.0, 3.0]
    style_table_header(arch_table.rows[0], col_w, ["Package / App", "Technology Stack", "Core Responsibility & Key Capabilities"])
    
    arch_rows = [
        ("apps/web\n(Frontend App)", "React 19, Vite 6, @xyflow/react (React Flow 12), Tailwind CSS 3, Dagre, Zustand, Lucide React, i18next", "Interactive infinite-canvas UI, custom node cards, auto-layout DAG engine, slide-over detail drawer, executive report modal, voice recording, multi-language switcher (EN, HI, KN)."),
        ("apps/api\n(Backend API)", "NestJS 10, Prisma ORM 6, PostgreSQL 16, Redis 7.2, Socket.io, BullMQ, Passport JWT, Swagger OpenAPI", "Secure REST API endpoints (/api/v1), WebSocket real-time presence and updates, RBAC authorization guards, audit trail capture, document metadata with BigInt serialization."),
        ("packages/shared\n(Shared Domain)", "TypeScript, Zod, Vitest", "Pure domain calculation engine (weighted average rollups, group blocker rules, overdue detectors), Zod input validation schemas, enterprise shared constants and type contracts."),
        ("Infrastructure &\nDevOps", "Docker Compose, PostgreSQL 16 Alpine, Redis 7.2 Alpine, MinIO S3 Object Storage", "Containerized local and production infrastructure, data persistence volumes, MinIO object store with console, GitHub Actions CI/CD pipeline verification.")
    ]
    for idx, row_data in enumerate(arch_rows):
        format_table_row(arch_table.rows[idx+1], col_w, row_data, is_even=(idx % 2 == 1))

    doc.add_paragraph()

    # SECTION 3: CORE CAPABILITIES & DEEP-DIVE FEATURES
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. In-Depth Functional Capabilities")
    h3_run.font.color.rgb = navy
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    # 3.1 Visual Canvas
    sub31 = doc.add_heading(level=2)
    sub31.add_run("3.1 Interactive Infinite-Canvas Studio (CanvasStudio.tsx)").font.color.rgb = accent_blue
    sub31.paragraph_format.space_before = Pt(8)
    sub31.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "The centerpiece of FlowTrack is the Canvas Studio, powered by React Flow (@xyflow/react) and custom state orchestrators. Key capabilities include:"
    )
    canvas_features = [
        ("Universal 4-Directional Handles", "Nodes feature interactive handles on Top, Bottom, Left, and Right, allowing users to draw dependency lines in any topological direction."),
        ("n8n-Style Floating Toolbar", "Hovering or selecting any node reveals a floating contextual action bar with one-click actions: Inspect Dossier, Add Next Step (+), Duplicate Task, Sub-Canvas Drill-Down, and Delete."),
        ("Sequence Edges with Inline Insertion (+)", "Custom SVG edges (FlowTrackEdge.tsx) feature an animated mid-point (+) button. Clicking this instantly splices a new intermediate task between two connected nodes without breaking existing connections."),
        ("Dual Auto-Layout Engine (Dagre)", "Includes one-click layout algorithms: 'Tidy Flow' (Left-to-Right sequence) and 'Tree' (Top-to-Bottom hierarchy). Floating, unlinked tasks are automatically organized into clean multi-column card grids."),
        ("Canvas MiniMap & Controls", "Provides smooth panning, infinite zoom, interactive mini-map locator, and fit-view keyboard shortcuts."),
        ("Right-Click Context Menu", "Right-clicking any node opens a desktop-style context menu for swift duplication, task inspection, and safe deletion.")
    ]
    for f_title, f_desc in canvas_features:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{f_title}: ")
        r.font.bold = True
        p.add_run(f_desc)

    # 3.2 Hierarchical Sub-Canvases
    sub32 = doc.add_heading(level=2)
    sub32.add_run("3.2 Sub-Canvas Drill-Down & Hierarchical Nesting").font.color.rgb = accent_blue
    sub32.paragraph_format.space_before = Pt(8)
    sub32.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "Complex initiatives quickly become unwieldy if thousands of tasks are plotted onto a single flat board. FlowTrack implements an infinite nested sub-canvas model:"
    )
    subcanvas_points = [
        ("Stage / Group Nodes", "High-level summary nodes that represent phases (e.g., 'Phase 1: Architecture', 'Phase 2: Execution'). Double-clicking or clicking 'Sub-Canvas' dives inside."),
        ("Breadcrumb Navigation Trail", "A dynamic breadcrumb bar (e.g., Root Canvas > Phase 1 > Database Setup) tracks navigation depth, allowing instant jumping between parent and child workspaces."),
        ("Direct Child Spawning", "Users can spawn interconnected sub-tasks directly beneath a parent task, automatically wiring them with a parent-to-child sequence edge.")
    ]
    for st, sd in subcanvas_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{st}: ").font.bold = True
        p.add_run(sd)

    # 3.3 Mathematical Rollup Engine
    sub33 = doc.add_heading(level=2)
    sub33.add_run("3.3 Deterministic Progress Rollup & Rule Engine (rollup.ts)").font.color.rgb = accent_blue
    sub33.paragraph_format.space_before = Pt(8)
    sub33.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "Rather than relying on team members to manually estimate overall stage completion, FlowTrack executes a mathematical rollup algorithm validated with 100% test coverage in packages/shared:"
    )

    create_callout_box(doc, [
        ("MATHEMATICAL ROLLUP FORMULATION", {"bold": True, "size": 10.5, "color": navy}),
        ("Weighted Average Progress = round( sum(Progress_i * Weight_i) / sum(Weight_i) )", {"bold": True, "size": 10, "color": accent_blue}),
        ("Rule 1 (Completion Guard): A parent node CANNOT be marked Done if any child node remains in an open or blocked state.", {"bold": False, "size": 9.5}),
        ("Rule 2 (Blocker Escalation): If any child node transitions to 'Blocked' / 'Waiting', the parent node automatically flags waiting count and suggests a Blocked status.", {"bold": False, "size": 9.5}),
        ("Rule 3 (Overdue Propagation): Past-due child milestones trigger an overdue badge on the parent group, alerting leadership immediately.", {"bold": False, "size": 9.5})
    ], border_color="2563EB", bg_color="F8FAFC")

    # 3.4 Multi-Tier RBAC Governance
    sub34 = doc.add_heading(level=2)
    sub34.add_run("3.4 Enterprise Role-Based Access Control (RBAC)").font.color.rgb = accent_blue
    sub34.paragraph_format.space_before = Pt(8)
    sub34.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "FlowTrack enforces strict enterprise permission boundaries defined in permissions.ts. User roles determine what actions are permitted:"
    )

    rbac_table = doc.add_table(rows=5, cols=4)
    rbac_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_rbac = [1.2, 1.3, 2.0, 2.0]
    style_table_header(rbac_table.rows[0], col_w_rbac, ["Role", "Badge Color", "Authorized Powers", "Enforced Restrictions"])

    rbac_data = [
        ("Super Admin", "Purple", "Edit canvas layout, auto-tidy DAG, user management, workspace administration, create/delete projects, full access to confidential files.", "None (Root Authority)"),
        ("Manager", "Blue", "Create/delete projects, assign tasks to anyone, approve/reject tasks (Thumbs Up/Down), approve deadline extensions, lock confidential nodes.", "Cannot alter canvas design layout (layout design alteration reserved for Super Admin)."),
        ("Team Lead", "Amber", "Assign tasks to their team members, create child subtasks, submit formal deadline extension requests, view project metrics.", "Cannot create/delete projects, cannot approve sign-offs, restricted from confidential locked documents."),
        ("Member", "Emerald", "Execute assigned tasks, log voice/photo remarks, update progress % slider, create sub-tasks for their assignments, request deadline extensions.", "Cannot reassign tasks, cannot mark approved without manager, restricted from confidential documents, cannot alter deadlines without formal justification.")
    ]
    for idx, rbac_row in enumerate(rbac_data):
        format_table_row(rbac_table.rows[idx+1], col_w_rbac, rbac_row, is_even=(idx % 2 == 1))

    doc.add_paragraph()

    # 3.5 Dual Dossier Inspection
    sub35 = doc.add_heading(level=2)
    sub35.add_run("3.5 Node Detail Drawer & Comprehensive Task Dossier").font.color.rgb = accent_blue
    sub35.paragraph_format.space_before = Pt(8)
    sub35.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "FlowTrack provides dual inspection mechanisms: a quick Slide-Over Drawer (NodeDetailDrawer.tsx) for fast edits during canvas navigation, and a Full Dossier Modal (NodeViewDetailsModal.tsx) for thorough auditing. Both share rich modules:"
    )

    dossier_features = [
        ("Remarks & Audio Notes", "Chronological collaborator timeline. Collaborators can type text, attach photographs, or record voice memos using browser Web Audio MediaRecorder. Features 5 quick status presets (e.g., 'Work In Progress', 'Waiting on Client')."),
        ("Manager Approval Gate", "Dedicated tab where managers perform official sign-offs: Thumbs Up (Approves and marks ready for delivery) or Thumbs Down (Rejects and forces a mandatory feedback remark)."),
        ("Document Vault & Confidential Locking", "Full file repository per node supporting versioning, file sizes, mime icons, inline image/PDF modal previews, and one-click confidential lock toggles."),
        ("Real-Time Audit Trail", "Displays every single lifecycle mutation (CREATE, UPDATE, STATUS_CHANGE) with before-and-after JSON diffs, timestamps, and collaborator attribution."),
        ("Dynamic Custom Fields", "Renders custom fields configured for the node's type (text, currency, date, selects, checkboxes) stored in Postgres JSONB."),
        ("Live Presence & Assignees", "Visual indicators displaying active online status for assignees, team avatars, and multi-collaborator assignment tags.")
    ]
    for d_title, d_desc in dossier_features:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{d_title}: ").font.bold = True
        p.add_run(d_desc)

    # 3.6 Deadline Extension Protocol
    sub36 = doc.add_heading(level=2)
    sub36.add_run("3.6 Audited Deadline Extension Protocol (DeadlineExtensionModal.tsx)").font.color.rgb = accent_blue
    sub36.paragraph_format.space_before = Pt(8)
    sub36.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "A common pitfall in task management is 'stealth rescheduling'—team members quietly pushing deadlines back without leadership awareness. FlowTrack eliminates this through a strict audit modal:"
    )
    doc.add_paragraph(
        "Whenever a user attempts to change an active task due date or when a task exceeds its timeline, the Deadline Extension Modal intercepts the action. The user must input a formal reason (e.g., 'Supplier shipment delayed by 48 hours') and can optionally attach an official justification document. This entry is permanently committed to the task's immutable Audit Log."
    )

    # 3.7 Executive Reporting & CSV Export
    sub37 = doc.add_heading(level=2)
    sub37.add_run("3.7 Executive Dossier & CSV Extraction (ProjectReportModal.tsx)").font.color.rgb = accent_blue
    sub37.paragraph_format.space_before = Pt(8)
    sub37.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "Executives and stakeholders require instant high-level visibility without navigating the full canvas. FlowTrack includes an Executive Audit Dossier:"
    )
    report_items = [
        ("KPI Summary Cards", "Live counts for Total Tasks, Completed Tasks, In-Progress Tasks, and Overdue Tasks, accompanied by a project-wide progress completion bar."),
        ("Task Breakdown Table", "Tabular extraction of all nodes displaying title, status badge, progress %, priority, countdown status, assignee, and attachment counts."),
        ("One-Click CSV Export", "Generates and downloads a clean, RFC-compliant CSV spreadsheet containing full project task metadata for Excel or BI tools."),
        ("Print-Ready Stylesheet", "Formatted with clean CSS print media queries for instant generation of executive PDF reports or paper dossiers.")
    ]
    for rt, rd in report_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{rt}: ").font.bold = True
        p.add_run(rd)

    # 3.8 Multilingual Support
    sub38 = doc.add_heading(level=2)
    sub38.add_run("3.8 Native Multilingual Architecture (i18next)").font.color.rgb = accent_blue
    sub38.paragraph_format.space_before = Pt(8)
    sub38.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "Recognizing diverse enterprise workforces (particularly in Indian development centers and multinational teams), FlowTrack comes pre-configured with first-class internationalization:"
    )
    i18n_table = doc.add_table(rows=4, cols=4)
    i18n_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_i18n = [1.5, 1.6, 1.7, 1.7]
    style_table_header(i18n_table.rows[0], col_w_i18n, ["Status Key", "English (en)", "Hindi (hi)", "Kannada (kn)"])
    i18n_data = [
        ("not_started", "Not Started", "शुरू नहीं हुआ", "ಪ್ರಾರಂಭವಾಗಿಲ್ಲ"),
        ("in_progress", "In Progress", "प्रगति पर है", "ಪ್ರಗತಿಯಲ್ಲಿದೆ"),
        ("waiting", "Waiting / Blocked", "अटका हुआ", "ಕಾಯುತ್ತಿದೆ")
    ]
    for idx, row in enumerate(i18n_data):
        format_table_row(i18n_table.rows[idx+1], col_w_i18n, row, is_even=(idx % 2 == 1))

    doc.add_paragraph()

    # SECTION 4: RELATIONAL DATA MODEL (PRISMA ORM)
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("4. Relational Database Schema & Entities")
    h4_run.font.color.rgb = navy
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "FlowTrack's backend is backed by PostgreSQL via Prisma ORM. The relational model is strictly typed and multi-tenant aware:"
    )

    schema_table = doc.add_table(rows=12, cols=3)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_schema = [1.5, 2.0, 3.0]
    style_table_header(schema_table.rows[0], col_w_schema, ["Entity / Model", "Core Database Columns", "Domain Purpose & Relational Connections"])

    schema_rows = [
        ("Workspace", "id, name, slug, logoUrl, themeSettings, timestamps", "Multi-tenant root container. Owns all users, projects, templates, and rules."),
        ("User", "id, email, phone, passwordHash, fullName, avatarUrl, locale", "System user entity supporting email/password and SMS OTP authentication."),
        ("Membership", "id, workspaceId, userId, role, projectPermissions", "Associates users with workspaces and defines their RBAC tier (ADMIN, MANAGER, etc.)."),
        ("Project", "id, workspaceId, name, description, status, settings, version", "Project container holding the directed workflow graphs and project-level chats."),
        ("NodeType", "id, workspaceId, key, name, icon, colour, isGroup, isSystem", "Defines node categories (Task, Stage, Approval Gate, Document Request, Milestone, Note)."),
        ("FieldDefinition", "id, nodeTypeId, fieldKey, label, type, required, options", "Configures dynamic custom fields (text, number, currency, select) per node type."),
        ("Node", "id, projectId, parentNodeId, nodeTypeId, statusId, ownerId, title, weight, progressPct, customData", "Core graph node. Supports parentNodeId for recursive sub-canvases and optimistic locking."),
        ("Edge", "id, projectId, sourceNodeId, targetNodeId, kind, label, metadata", "Directed link between nodes with kind: 'sequence', 'dependency', or 'related'."),
        ("Remark", "id, nodeId, authorId, body, voiceUrl, photoUrl, createdAt", "Collaborator comments, voice memos, and photo attachments attached to a node."),
        ("Document", "id, nodeId, uploadedBy, name, mimeType, fileSize, storageKey, scanStatus", "File attachments stored in MinIO with mime validation and BigInt size serialization."),
        ("AuditLog", "id, workspaceId, projectId, entityName, entityId, actorId, action, oldValues, newValues", "Immutable historical log tracking all mutations (CREATE, UPDATE, DELETE, STATUS_CHANGE).")
    ]
    for idx, s_row in enumerate(schema_rows):
        format_table_row(schema_table.rows[idx+1], col_w_schema, s_row, is_even=(idx % 2 == 1))

    doc.add_paragraph()

    # SECTION 5: HOW EVERYTHING CONNECTS & USER WORKFLOW
    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run("5. End-to-End Operational Lifecycle")
    h5_run.font.color.rgb = navy
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "To understand how FlowTrack functions during daily operations, consider this end-to-end lifecycle walkthrough:"
    )

    steps = [
        ("Step 1: Workspace & Project Initialization", "An Administrator logs in and creates a new project (e.g., 'Commercial Complex Construction' or 'SaaS V2 Launch'). The system seeds default status sets and node types."),
        ("Step 2: Workflow Canvas Construction", "The project manager maps out the process on the infinite canvas. They drop Stage nodes for main phases, connect them with sequence edges, and use the (+) edge insert tool to inject approval gates."),
        ("Step 3: Sub-Canvas Decomposition", "Managers double-click a Stage node to enter its isolated sub-canvas. Here, granular tasks are created, assigned to team leads, given estimated weights (e.g., 2.5x impact), and given concrete due dates."),
        ("Step 4: Daily Execution & Remark Logging", "Assigned team members open their tasks. They adjust progress % sliders, log voice memos describing site inspections, and upload PDFs or images. If an impediment arises, they mark the task 'Waiting / Blocked'."),
        ("Step 5: Automated Progress & Status Rollup", "The shared calculation engine computes the Stage's progress instantly. If any child task is blocked, the stage immediately signals a warning badge without manual intervention."),
        ("Step 6: Manager Review & Sign-Off", "Upon reaching 100%, the manager opens the Approval Gate tab in the dossier, reviews attached test reports or legal documents, and issues a formal 'Thumbs Up' sign-off."),
        ("Step 7: Executive Reporting & Export", "Executive leadership opens the Project Report Modal to review overall delivery health, print the dossier for board meetings, or export CSV datasets for company-wide analytics.")
    ]
    for s_title, s_desc in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{s_title}: ").font.bold = True
        p.add_run(s_desc)

    doc.add_paragraph()

    # SECTION 6: QUICK START & VERIFICATION GUIDE
    h6 = doc.add_heading(level=1)
    h6_run = h6.add_run("6. Local Development & Deployment Guide")
    h6_run.font.color.rgb = navy
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The project is ready for local execution and testing. Follow these steps to launch all services:"
    )

    setup_table = doc.add_table(rows=6, cols=3)
    setup_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w_setup = [1.2, 2.5, 2.8]
    style_table_header(setup_table.rows[0], col_w_setup, ["Step", "Command / Endpoint", "Description"])
    setup_data = [
        ("1. Dependencies", "pnpm install", "Installs all dependencies across apps/web, apps/api, and packages/shared."),
        ("2. Shared Lib", "pnpm --filter @flowtrack/shared build", "Compiles domain Zod schemas and calculation engines to dist/."),
        ("3. Database", "cd apps/api && npx prisma db push && npx ts-node prisma/seed.ts", "Deploys PostgreSQL schema and seeds default users, roles, and status sets."),
        ("4. Start API", "pnpm --filter @flowtrack/api start:dev", "Launches NestJS backend at http://localhost:3001 with Swagger docs at /api/docs."),
        ("5. Start Web", "pnpm --filter @flowtrack/web dev", "Launches Vite React app at http://localhost:5173.")
    ]
    for idx, s_row in enumerate(setup_data):
        format_table_row(setup_table.rows[idx+1], col_w_setup, s_row, is_even=(idx % 2 == 1))

    doc.add_paragraph()

    # Default Credentials Box
    create_callout_box(doc, [
        ("DEFAULT SEED CREDENTIALS FOR TESTING", {"bold": True, "size": 10.5, "color": navy}),
        ("Super Admin: admin@flowtrack.dev | Password: Password123!", {"bold": False, "size": 9.5}),
        ("Project Manager: suresh@flowtrack.dev | Password: Password123!", {"bold": False, "size": 9.5}),
        ("Team Member: ramesh@flowtrack.dev | Password: Password123!", {"bold": False, "size": 9.5}),
        ("Dev SMS OTP Bypass: Enter phone number and use code '123456'", {"bold": False, "size": 9.5})
    ], border_color="10B981", bg_color="ECFDF5")

    # Output file path
    output_path = r"c:\Users\User\Desktop\Project-managment-app\FlowTrack_Project_Overview_and_Agenda.docx"
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

if __name__ == "__main__":
    generate_report()
